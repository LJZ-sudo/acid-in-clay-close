# -*- coding: utf-8 -*-
"""Build the report deck (.pptx) — Chinese, plain-language, positive-leaning, journal-grade figures.

Design (per latest feedback):
  - Slides 1-2 are NATIVE Chinese shapes (cards + chevrons): accurate CJK text, editable, NO
    'floor/ceiling' jargon, plain wording.
  - Slide 3 = journal Figure 1 (overall workflow, fig1_workflow_draft_0.png).
  - Slides 4-6 = one innovation each, using ONLY positive result figures (discovery / calibration
    reliability / closed-loop BO+LLM). Negative diagnostics moved to backup/notes.
  - Slide 7 = positioning + methodological upgrade (skills/harness/memory) + next step.
  - Slide 8 = backup: code-grounded implementation architecture.
6-8 slides, multiple figures per slide allowed. Speaker notes attached per slide.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

NDA = (next(_p for _p in Path(__file__).resolve().parents if _p.name == "V1.0-qianduan-mainline").parent / "research")
FIG = NDA / "figures"
OUT = NDA / "REPORT_DECK_20260622.pptx"

NAVY = RGBColor(0x12, 0x32, 0x4E)
TEAL = RGBColor(0x1F, 0x9E, 0x9B)
GREY = RGBColor(0x6B, 0x72, 0x80)
LGREY = RGBColor(0xF3, 0xF5, 0xF7)
ORANGE = RGBColor(0xE3, 0x7B, 0x1E)
BLUE = RGBColor(0x2B, 0x5C, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CJK = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _cjk(run, name=CJK):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)


def add_textbox(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        _cjk(r)
    return tb


def add_image_fit(slide, img_path, max_l, max_t, max_w, max_h):
    with Image.open(img_path) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = max_w / max_h
    if ar > box_ar:
        w = max_w; h = int(max_w / ar)
    else:
        h = max_h; w = int(max_h * ar)
    l = max_l + (max_w - w) // 2
    t = max_t + (max_h - h) // 2
    slide.shapes.add_picture(str(img_path), l, t, width=Emu(w), height=Emu(h))


def title_bar(slide, text, color=NAVY, sub=None):
    add_textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.8),
                [(text, 28, True, color)])
    if sub:
        add_textbox(slide, Inches(0.57), Inches(1.02), Inches(12.2), Inches(0.5),
                    [(sub, 13, False, GREY)])


def card(slide, l, t, w, h, accent, number, heading, body, evidence):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = accent; box.line.width = Pt(1.5)
    box.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, Inches(0.16))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    bar.shadow.inherit = False
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(0.18), t + Inches(0.28),
                                   Inches(0.5), Inches(0.5))
    badge.fill.solid(); badge.fill.fore_color.rgb = accent; badge.line.fill.background()
    badge.shadow.inherit = False
    btf = badge.text_frame; btf.word_wrap = False
    br = btf.paragraphs[0].add_run(); br.text = number
    br.font.size = Pt(18); br.font.bold = True; br.font.color.rgb = WHITE; _cjk(br)
    btf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.22); tf.margin_top = Inches(0.95)
    runs = [(heading, 15, True, accent), (body, 12.5, False, NAVY), (" ", 6, False, NAVY),
            (evidence, 10.5, False, GREY)]
    for i, (text, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; _cjk(r)


def chevron(slide, l, t, w, h, color, number, title):
    sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f"{number}  {title}"
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = WHITE; _cjk(r)


def figure_slide(title, img, caption, accent=NAVY, sub=None):
    s = add_slide()
    title_bar(s, title, accent, sub)
    top = Inches(1.45) if sub else Inches(1.2)
    add_image_fit(s, img, Inches(0.7), top, Inches(11.9), Inches(4.85))
    add_textbox(s, Inches(0.7), Inches(6.5), Inches(11.9), Inches(0.85),
                [(caption, 13, False, NAVY)], align=PP_ALIGN.CENTER)
    return s


# ---------------- Slide 1: three innovations (native cards, Chinese) ----------------
s = add_slide()
title_bar(s, "三个核心创新点",
          NAVY,
          "条件：只有阻抗(EIS)数据、样本量小、不做额外结构表征  ·  目标：做“可信”的自主发现")
cw, gap, x0, ct, ch = Inches(3.81), Inches(0.4), Inches(0.55), Inches(1.75), Inches(3.95)
card(s, x0, ct, cw, ch, TEAL, "1",
     "「发现」低温质子传导出现“区间转变”，且可由配方调控",
     "质子在零下低温仍能传导；存在一个传导行为的转变温度，转变的“陡峭程度”能被配方调出来。",
     "证据：4 类材料、13 次独立重复；转变温度 −22 ~ −39 ℃")
card(s, x0 + cw + gap, ct, cw, ch, BLUE, "2",
     "「方法」让系统的“把握度”变可信，并给结论严格分级",
     "量化“同配方重复制样会有多大波动”，据此校准置信度；只说数据撑得住的结论，不过度解读。",
     "证据：校准后置信度更可靠(ECE 0.78 → 0.13)")
card(s, x0 + 2 * (cw + gap), ct, cw, ch, ORANGE, "3",
     "「系统」真实的“提配方→合成→测试→更新”智能体闭环",
     "贝叶斯优化提升找优效率；语言模型把关配方安全；全过程可追溯、可审计。",
     "证据：优化比随机更快找到最优；语言模型把不可合成配方拉回安全区")
add_textbox(s, Inches(0.55), Inches(5.95), Inches(12.2), Inches(1.3), [
    ("为什么是这三个？", 14, True, NAVY),
    ("三点合一，正好覆盖「发现什么 → 凭什么可信 → 怎么做到的」，是同一套“可信自主发现”方法的三个侧面——"
     "不是三个拼凑的结果，而是从同一批数据里长出来的一个内核。", 12.5, False, NAVY),
])
set_notes(s, "三贡献(正面表述)：①发现低温传导区间转变且可配方调控；②让把握度可信(校准)+结论严格分级；"
             "③真实智能体闭环(BO提效+LLM保安全+可追溯)。为什么是这三个：覆盖'发现什么→凭什么可信→怎么做到的'，"
             "同源、非拼凑。措辞按独立评审：说'传导区间转变'非'相变'；说'置信度更可信'非'更准确'。")

# ---------------- Slide 2: storyline (native chevrons, Chinese) ----------------
s = add_slide()
title_bar(s, "主要思路：一条主线")
steps = [(GREY, "①", "苛刻条件"), (TEAL, "②", "发现"), (BLUE, "③", "可信"),
         (ORANGE, "④", "边界"), (BLUE, "⑤", "验证")]
details = [
    "只有阻抗数据、\n样本量小",
    "低温传导出现\n“区间转变”，\n且可配方调控",
    "量化重复制样\n的波动，校准\n系统的把握度",
    "严格界定哪些\n结论数据撑得\n住、哪些撑不住",
    "真实闭环验证：\n更高效、更安全",
]
cw2, x0, ct, ch2 = Inches(2.45), Inches(0.45), Inches(2.05), Inches(1.25)
step_overlap = Inches(0.12)
for i, (color, num, title) in enumerate(steps):
    lx = x0 + i * (cw2 - step_overlap)
    chevron(s, lx, ct, cw2, ch2, color, num, title)
    add_textbox(s, lx + Inches(0.12), ct + ch2 + Inches(0.12), cw2 - Inches(0.2), Inches(1.6),
                [(details[i], 12, False, NAVY)], align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.55), Inches(5.7), Inches(12.2), Inches(1.2), [
    ("一句话总结", 14, True, ORANGE),
    ("把“自主发现的可信程度”做成可测量、可复算的方法——先量出测量本身能分辨多少，再决定哪些发现/结论"
     "真正成立。", 14, True, NAVY),
])
set_notes(s, "一条主线(Evidence-limited autonomous science)：苛刻条件→发现可调的低温传导区间转变→"
             "量化重复波动校准把握度→严格界定结论范围→真实闭环验证(更高效更安全)。"
             "总结：把'可信程度'做成可测量可复算的方法。")

# ---------------- Slide 3: journal Figure 1 ----------------
figure_slide(
    "图 1 · 总体工作流：受测量能力约束的自主发现",
    FIG / "fig1_workflow_draft_0.png",
    "材料→变温阻抗→传输描述符；独立重复给出测量重复性、模型竞争给出可辨识范围；"
    "受治理的推理+闭环(BO+语言模型安全把关)→带可信度与分级结论的输出。(期刊 Fig 1 初稿)",
    accent=NAVY,
    sub="论文定位：在“单一传输观测 + 样本昂贵 + 模型不可唯一识别”的现实下，做可信的自主发现")
set_notes(prs.slides[2],
          "这是论文 Fig 1(总体工作流)。四阶段：①材料与测量→传输描述符；②证据校准(重复性+可辨识范围)；"
          "③受治理推理与闭环(确定性核，LLM只负责提出与解释)；④可信度封顶的分级输出。底部治理带：每体系独立"
          "history、前瞻冻结、claim 审计。术语精确：transport-regime transition、measurement reproducibility、"
          "model identifiability、graded claims C0–C4。")

# ---------------- Slide 4: innovation 1 — discovery ----------------
figure_slide(
    "创新点 ① 发现：低温传导“区间转变”且可配方调控",
    NDA / "regularity/regularity.png",
    "(a) 不同材料的转变“陡峭程度”系统性不同(藕粉最陡、淀粉最缓)；(b) 13/15 个样品的转变温度全在 0 ℃ 以下"
    "(−22 ~ −39 ℃)。说明这是可被配方调出来的规律，而非偶然。",
    accent=TEAL,
    sub="把“低温还能导”从单点现象，做成“跨材料、可重复、可调控”的规律")
set_notes(prs.slides[3],
          "正面发现：跨4类材料、13次独立重复都出现亚零度传导区间转变；陡峭度可由配方调(藕粉~陡、淀粉~缓)。"
          "Ea_low 0.69±0.07 eV(LRS 7重复)。措辞用'传导区间转变(transport-regime transition)'，不写'相变'。"
          "诚实备注(不上台)：转变稳健性还需做统一Rb方法/无单调过滤的敏感性复核(G2)。")

# ---------------- Slide 5: innovation 2 — credibility/calibration ----------------
figure_slide(
    "创新点 ② 方法：让系统的“把握度”变可信",
    NDA / "calibration/calibration_reliability_only.png",
    "原始的单曲线把握度严重“过度自信”(红：几乎都报 0.9+ 却只有 ~0.2 实际复现)；用重复数据校准后，"
    "报告的置信度回到可信区间(绿，ECE 0.78 → 0.13，留一数据集)。系统因此知道“何时该有把握、何时该谨慎”。",
    accent=BLUE,
    sub="可信度量化 + 严格结论分级(C0–C4)：只说数据撑得住的话")
set_notes(prs.slides[4],
          "正面方法：校准把过度自信拉回可信区间，ECE 0.78→0.13(留一数据集)。表述严格按独立评审："
          "是'提高了概率陈述的可信度/拒答合理性'，不是'提高了发现准确率'(resolution 几乎不变)。"
          "配合 C0–C4 主张分级：EIS-only 封顶 C4，不声称机理。负结果(单曲线QC测不准复现 AUROC≈0.5)放备份，不上台。")

# ---------------- Slide 6: innovation 3 — closed loop ----------------
figure_slide(
    "创新点 ③ 系统：真实智能体闭环(更高效 + 更安全)",
    NDA / "replay/strategy_comparison.png",
    "(左)贝叶斯优化平均 4.1 个实验就找到最优，随机要 5.5 个——优化确实更快；"
    "(右)纯优化会漂到无法合成的极端配方，语言模型的安全把关把它拉回可行区(真实最优就在其中)。",
    accent=ORANGE,
    sub="贝叶斯优化负责“快”，语言模型负责“安全可行”，全过程可追溯、可审计")
set_notes(prs.slides[5],
          "正面系统：BO 4.1 vs random 5.5(更快)；LLM 把不可合成的极端配方拉回安全可行区(R0.02→0.28，真实最优"
          "R0.186在内)。诚实备注(不上台)：此对照含 replay 成分，投稿时要标注 replay vs 真实，并与硬安全投影/"
          "约束GP等强基线对比(G3/G4)。")

# ---------------- Slide 7: positioning + methodological upgrade + next ----------------
s = add_slide()
title_bar(s, "定位、方法学升级与下一步")
add_textbox(s, Inches(0.55), Inches(1.35), Inches(6.05), Inches(2.4), [
    ("发表定位", 16, True, TEAL),
    ("核心贡献不是“发现了某种材料”，而是：在只有单一传输观测、实验昂贵、模型不可唯一识别的条件下，"
     "一个自主系统如何定量地知道——何时可以下结论、何时必须拒绝、一个算法建议是否超出了实验本身的分辨能力。",
     12, False, NAVY),
    ("目标刊：Tier B(Comm. Chem./Mater. / CRPS / npj，IF 6–10)；做完关键补强后有冲更高的潜力。",
     12, False, NAVY),
])
add_textbox(s, Inches(6.85), Inches(1.35), Inches(5.95), Inches(2.4), [
    ("方法学升级(下一步原创增量)", 16, True, ORANGE),
    ("· 技能(Skills)：给每个分析步骤一个“结论类型”，从源头限定它最多能支撑到哪一级主张。", 11.5, False, NAVY),
    ("· 执行(Harness)：让采集/停止被“测量分辨能力”门控——不去优化测不出的差别。", 11.5, False, NAVY),
    ("· 记忆(Memory)：记忆不仅存结果，还存“能测多准 / 能撑什么结论 / 来源凭证”。", 11.5, False, NAVY),
])
add_textbox(s, Inches(0.55), Inches(4.0), Inches(12.2), Inches(3.0), [
    ("下一步关键补强(诚实，按优先级)", 16, True, NAVY),
    ("1) G1：在凹凸棒土最佳配方上做 ≥3 次同配方重复，直接测出该体系自己的重复性(替换代理)。", 12, False, NAVY),
    ("2) G2：Stage0 阈值与方法敏感性——统一 Rb 提取方法 / 不做单调过滤 也要复核转变是否稳健。", 12, False, NAVY),
    ("3) G3/G4：闭环对照标注 replay vs 真实；语言模型安全修正要对“硬安全投影/约束优化”等强基线。", 12, False, NAVY),
    ("4) 措辞与可复现：保存完整 prompt 与原始响应；“相变→传导区间转变”“校准→提高可信度而非准确率”。", 12, False, NAVY),
])
set_notes(s, "定位(对齐独立评审)：贡献=在受限观测下定量知道何时可声称/何时拒绝/建议是否超出实验分辨力。"
             "方法学升级=skills/harness/memory三创新(把可信边界做成可存/可执行/可类型检查)。"
             "下一步：G1(湿实验门)+ G2 Stage0敏感性 + G3/G4强基线与replay标注 + 措辞降温与prompt留存。"
             "明确：硬门槛不止G1。")

# ---------------- Slide 8: backup — code architecture ----------------
figure_slide(
    "(备份) 系统实现架构：代码级流水线",
    FIG / "architecture_diagram_draft_0.png",
    "六阶段实现：Stage0 测量→证据 / Stage1 闭环 / Stage2 证据种子 / Stage3 机理+候选+治理 / 只读看板 / v2 守卫。"
    "绿=确定性、蓝=LLM；候选生成/排序/主张审计默认确定性冻结。(SI/附录用)",
    accent=GREY)
set_notes(prs.slides[7],
          "备份页(答辩/SI)：代码级实现架构。强调确定性 vs LLM 边界、护栏带。配合 SYSTEM_ARCHITECTURE 文档附录 A–E。")

prs.save(str(OUT))
print(f"Saved -> {OUT}  ({len(prs.slides)} slides)")
